#!/usr/bin/env python3
"""html-to-pptx.py — optional PPTX snapshot export for a feishu-deck-h5 deck.

Turns a confirmed HTML deck into a PPTX where each slide is a faithful full-bleed
image of the rendered HTML page (present mode, 1920×1080 design box, no chrome).
Reuses the skill's own `deck-json/shoot.py` (Playwright) for screenshots and
assembles the PPTX with python-pptx — no external pipeline dependency.

TRADEOFF — read before using:
  This is a SNAPSHOT pptx. It is visually identical to the HTML (the exact thing
  a user/agent iterated to quality), but **text is NOT editable** — each slide is
  one picture. Animations / transitions / iframes are flattened to their resting
  state. For an editable native PPTX (real shapes/text), author in a dedicated
  PowerPoint / Keynote workflow; this exporter is the "faithful PPT copy of my
  HTML deck" option for when sharing a .pptx matters more than editability.

Usage:
    python3 assets/html-to-pptx.py <index.html | run-dir> [--out deck.pptx] [--pages 1,2,cover]
"""
import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL_ROOT = HERE.parent
SHOOT = SKILL_ROOT / "deck-json" / "shoot.py"


def main() -> int:
    ap = argparse.ArgumentParser(description="Snapshot-export a feishu-deck-h5 HTML deck to PPTX (image per slide).")
    ap.add_argument("target", help="index.html path OR run/output dir holding it")
    ap.add_argument("--out", help="output .pptx path (default: beside index.html, <stem>.pptx)")
    ap.add_argument("--pages", help="comma slide indices/keys (default: all), passed to shoot.py")
    args = ap.parse_args()

    if not SHOOT.is_file():
        print(f"[error] shoot.py missing: {SHOOT}", file=sys.stderr)
        return 2
    try:
        from pptx import Presentation  # noqa: F401
    except ImportError:
        print("[error] python-pptx not installed: `pip install python-pptx`", file=sys.stderr)
        return 2

    target = Path(args.target).expanduser().resolve()
    index = target if target.is_file() else target / "index.html"
    if not index.is_file():
        print(f"[error] no index.html at {index}", file=sys.stderr)
        return 2
    out_pptx = Path(args.out).resolve() if args.out else index.with_suffix(".pptx")

    with tempfile.TemporaryDirectory(prefix="html_pptx_") as tmp:
        shoot_cmd = [sys.executable, str(SHOOT), str(index), "--out", tmp]
        if args.pages:
            shoot_cmd += ["--pages", args.pages]
        print(f"[shoot] rendering present-mode slides → PNG (Playwright) …")
        rc = subprocess.run(shoot_cmd).returncode
        if rc != 0:
            print(f"[error] shoot.py failed (exit {rc}); is Playwright + chromium installed?", file=sys.stderr)
            print("        pip install playwright && python -m playwright install chromium", file=sys.stderr)
            return rc
        pngs = sorted(Path(tmp).glob("p*.png"))  # p{idx:02d}-{key}.png → idx-padded, sorts in slide order
        if not pngs:
            print("[error] shoot.py produced no PNGs", file=sys.stderr)
            return 1
        print(f"[shoot] {len(pngs)} slide(s) captured")

        from pptx import Presentation
        from pptx.util import Emu
        prs = Presentation()
        prs.slide_width = Emu(12192000)   # 16:9 — 13.333"
        prs.slide_height = Emu(6858000)   # 7.5"
        blank = prs.slide_layouts[6]
        for png in pngs:
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
        prs.save(str(out_pptx))

    sz = out_pptx.stat().st_size / (1024 * 1024)
    print(f"\n[done] PPTX snapshot → {out_pptx}  ({sz:.1f} MB, {len(pngs)} slides)")
    print("       snapshot mode: visually faithful to the HTML; text is NOT editable.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
